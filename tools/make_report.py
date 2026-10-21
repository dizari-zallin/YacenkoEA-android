# -*- coding: utf-8 -*-
"""Builds a Times New Roman lab report and inserts captured screenshots."""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips

ROOT = Path(r"C:\Users\Shmurdyakich\Desktop\androidALL")
SHOTS = ROOT / "screenshots"
OUT = ROOT / "Отчет_практические_работы.docx"

LABS = [
    {
        "n": 0,
        "title": "Вводное занятие. Установка и настройка Android Studio",
        "goal": "Изучить порядок установки и настройки Android Studio на персональном компьютере.",
        "body": [
            "На компьютере установлена Android Studio версии 2026.1.4. Вместе со средой установлен комплект разработки: JDK из состава Android Studio (OpenJDK 25), Android SDK с платформой API 37.0 и средствами сборки Build-Tools 36.0.0. Каталог SDK расположен в локальном профиле пользователя.",
            "Эмулятор устройств входит в установленный SDK. Для выполнения практических работ создано виртуальное устройство и на нём проверены учебные приложения. Проекты сохранены в каталоге projects и открываются в Android Studio как отдельные приложения на языке Java.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с порядком установки и настройки Android Studio на персональном компьютере.",
    },
    {
        "n": 1,
        "title": "Практическая работа № 1. Интерфейс Android Studio",
        "goal": "Знакомство с инструментами разработки Android-приложений, с интерфейсом Android Studio, содержанием проекта и основными элементами управления.",
        "body": [
            "Создан проект YacenkoEA_PR1 по шаблону Empty Views Activity. Язык исходного кода — Java, минимальная версия Android — API 24. Модуль app содержит манифест, класс MainActivity и ресурсы. Файл activity_main.xml описывает экран с надписью Hello World! в контейнере ConstraintLayout.",
            "В дереве проекта на вкладке Android видны каталоги manifest, java и res, а также сценарии Gradle. Сборка debug-варианта выполнена плагином Android Gradle Plugin. После установки пакета на эмулятор приложение показывает заголовок Яценко PR1 и текст Hello World!.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с интерфейсом Android Studio, структурой проекта и запуском приложения в эмуляторе.",
    },
    {
        "n": 2,
        "title": "Практическая работа № 2. Поля ввода",
        "goal": "Знакомство с основными элементами в редакторе XML Android Studio. Создание приложения с использованием элементов ImageButton, TextView, EditText и обработчика нажатия на кнопку.",
        "body": [
            "В проекте YacenkoEA_PR2 на экране размещены поле EditText, кнопка, TextView и ImageButton. Картинка при старте скрыта атрибутом android:visibility=\"invisible\". Связи элементов заданы ограничениями ConstraintLayout, поэтому положение виджетов сохраняется при запуске.",
            "В MainActivity ссылки на виджеты получены методом findViewById. Объект View.OnClickListener назначен кнопке через setOnClickListener. По нажатию в текстовое поле выводится строка из EditText, а если строка пустая — текст «Все ок)». ImageButton становится видимым.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с полями ввода, кнопкой, текстовым полем, ImageButton и обработкой нажатия.",
    },
    {
        "n": 3,
        "title": "Практическая работа № 3. Обработчики событий",
        "goal": "Изучить основные способы обработки активации кнопки. Написать скрипт для подсчёта количества нажатий на кнопку.",
        "body": [
            "В проекте YacenkoEA_PR3 на экране находятся TextView с начальным значением 0 и кнопка «Нажми меня». Счётчик объявлен полем анонимного класса View.OnClickListener, поэтому значение сохраняется между нажатиями.",
            "В методе onClick счётчик увеличивается на единицу и записывается в текстовое поле через String.valueOf. Каждое нажатие обновляет число на экране.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с обработчиком нажатия кнопки и подсчётом количества нажатий.",
    },
    {
        "n": 4,
        "title": "Практическая работа № 4. Строковые и цветовые ресурсы",
        "goal": "Изучение правил создания строковых и цветовых ресурсов проекта, создание приложения с использованием этих ресурсов.",
        "body": [
            "В проекте YacenkoEA_PR4 имя приложения и фраза Hello Android! вынесены в res/values/strings.xml. Заголовки страниц записаны в отдельный файл headers.xml, цвета текста и фона — в my_colors.xml. В разметке текст и оформление задаются ссылками @string и @color.",
            "В MainActivity строка читается методом getString, цвет — методом getColor. Второй TextView получает текст и цвета уже из кода, поэтому один и тот же ресурс используется и в XML, и в Java.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с созданием и применением строковых и цветовых ресурсов.",
    },
    {
        "n": 5,
        "title": "Практическая работа № 5. Переключение между страницами",
        "goal": "Создать приложение из двух окон с возможностью переключения активного окна по нажатию кнопки.",
        "body": [
            "В проекте YacenkoEA_PR5 добавлены разметка activity_about.xml и класс AboutActivity. На первом экране отображается текст «Первое окно», на втором — «Второе окно». AboutActivity объявлена в AndroidManifest.xml.",
            "Переход выполняется объектом Intent: из MainActivity указывается класс AboutActivity, из AboutActivity — класс MainActivity. Запуск второго экрана выполняет метод startActivity.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с переключением между двумя активностями по нажатию кнопки.",
    },
    {
        "n": 6,
        "title": "Практическая работа № 6. Передача управления между страницами",
        "goal": "Создать приложение из двух окон с возможностью передачи данных от одной страницы к другой в одну сторону.",
        "body": [
            "В проекте YacenkoEA_PR6 описан абстрактный класс Human с полями имени, фамилии и возраста и класс-наследник Student с полем группы. Оба класса поддерживают сериализацию, поэтому объект можно положить в Intent.",
            "На первом экране четыре поля ввода и кнопка «Передать данные». По нажатию создаётся Student и передаётся во SecondActivity методом putExtra. Второй экран читает объект из Bundle и показывает имя, группу и возраст.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с односторонней передачей объекта между активностями через Intent.",
    },
    {
        "n": 7,
        "title": "Практическая работа № 7. Обмен данными между страницами",
        "goal": "Создать приложение из двух окон с возможностью передачи данных между страницами в обе стороны.",
        "body": [
            "В проекте YacenkoEA_PR7 первая активность регистрирует обработчик результата методом registerForActivityResult и контрактом StartActivityForResult. Кнопка «Заполнить форму» открывает второй экран.",
            "На втором экране вводится имя. Кнопка «Сохранить» упаковывает строку в Intent, возвращает результат RESULT_OK и закрывает экран методом finish. Первая активность получает имя в onActivityResult и выводит его в TextView.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с двусторонним обменом данными между активностями через Activity Result API.",
    },
    {
        "n": 8,
        "title": "Практическая работа № 8. Работа с изображениями",
        "goal": "Создать приложение с возможностью выбора картинки из библиотеки и её масштабирования.",
        "body": [
            "В проекте YacenkoEA_PR8 изображение выводится в ImageView. Кнопка «Галерея» открывает системный выбор изображения действием ACTION_PICK. В onActivityResult поток выбранного файла декодируется в Bitmap и передаётся в ImageView.",
            "Отдельные кнопки меняют android-свойство ScaleType: CENTER, CENTER_CROP, CENTER_INSIDE, FIT_CENTER и FIT_XY. Текущий режим подписывается над изображением, поэтому различие режимов масштабирования видно на экране.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с загрузкой изображения и режимами масштабирования ImageView.",
    },
    {
        "n": 9,
        "title": "Практическая работа № 9. Создание приложения с аудиоплеером",
        "goal": "Создать приложение с возможностью выбора аудио из библиотеки и управления его воспроизведением.",
        "body": [
            "В проекте YacenkoEA_PR9 аудиофайл music.wav помещён в каталог res/raw. Объект MediaPlayer создаётся методом create по идентификатору ресурса. Кнопки Play, Pause и Stop вызывают start, pause и stop. После остановки плеер подготавливается снова и перематывается в начало.",
            "Доступность кнопок зависит от состояния воспроизведения: до старта активна только кнопка Play. SeekBar связан с AudioManager и меняет громкость потока STREAM_MUSIC. По окончании ролика срабатывает OnCompletionListener и воспроизведение останавливается.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с воспроизведением аудио через MediaPlayer и регулировкой громкости.",
    },
    {
        "n": 10,
        "title": "Практическая работа № 10. Создание приложения с видеоплеером",
        "goal": "Создать приложение с возможностью выбора видео из библиотеки и управления его воспроизведением.",
        "body": [
            "В проекте YacenkoEA_PR10 на экране размещён VideoView и кнопки Play, Pause и Stop. Ролик clip.mp4 хранится в res/raw и подключается через URI вида android.resource. Кнопки вызывают start, pause и stopPlayback.",
            "Отдельная кнопка «Из интернета» передаёт в setVideoPath адрес ролика, указанный в методических указаниях. Для этого в манифест добавлено разрешение INTERNET и разрешён ясный HTTP-трафик. Основная проверка выполнена на локальном файле.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с воспроизведением видео через VideoView и управлением пуском, паузой и остановкой.",
    },
    {
        "n": 11,
        "title": "Практическая работа № 11. Создание приложения с кадровой анимацией",
        "goal": "Создать приложение с возможностью воспроизведения покадровой анимации.",
        "body": [
            "В проекте YacenkoEA_PR11 в каталог drawable помещены восемь кадров и файл animation.xml со списком animation-list. У каждого кадра задана длительность 180 мс. Атрибут oneshot=\"false\" включает циклический повтор.",
            "ImageView получает эту анимацию как фон. Объект AnimationDrawable запускается кнопкой «Старт» и останавливается кнопкой «Стоп». Пока анимация идёт, кадры последовательно сменяют друг друга.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с покадровой анимацией на основе AnimationDrawable.",
    },
    {
        "n": 12,
        "title": "Практическая работа № 12. Создание приложения с анимацией свойств",
        "goal": "Создать приложение с возможностью воспроизведения анимации свойств.",
        "body": [
            "В проекте YacenkoEA_PR12 в каталоге res/anim созданы два описания вращения: rotate_center.xml с центром 50% на 50% и rotate_corner.xml с центром в левом верхнем углу. Оба файла используют элемент rotate и линейный интерполятор.",
            "Кнопки «Вокруг центра» и «Вокруг угла» загружают соответствующую анимацию через AnimationUtils.loadAnimation и запускают её у ImageView методом startAnimation. Одно и то же изображение вращается относительно разных точек.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с анимацией вращения объекта относительно центра и угла.",
    },
    {
        "n": 13,
        "title": "Практическая работа № 13. Анимация геометрических объектов",
        "goal": "Создать приложение с возможностью анимации геометрических объектов в двумерном или трёхмерном пространстве.",
        "body": [
            "В проекте YacenkoEA_PR13 небо, трава и солнце заданы фигурами shape в каталоге drawable: прямоугольные градиенты и овальное солнце с радиальным градиентом. Фигуры выведены тремя ImageView внутри ConstraintLayout. Подписи связаны со строковыми ресурсами.",
            "Файл res/anim/sun_rise.xml объединяет масштаб, вертикальный перенос и изменение прозрачности. Длительность анимации — 5 секунд, fillAfter сохраняет конечное положение. Кнопка «Восход» запускает эту анимацию у изображения солнца.",
        ],
        "conclusion": "В результате выполнения работы было произведено ознакомление с анимацией геометрических фигур: масштабом, перемещением и прозрачностью.",
    },
    {
        "n": 14,
        "title": "Практическая работа № 14. Индивидуальное задание. Приложение «Авто»",
        "goal": "Создать приложение «Авто» по задаче 6: показать список автомобилей с характеристиками, добавить новый автомобиль и отредактировать его данные.",
        "body": [
            "В проекте YacenkoEA_PR14 выполнена задача 6 методички. В локальной базе SQLite хранятся 12 автомобилей трёх производителей. У Toyota две марки — Camry и Corolla, у BMW две марки — X5 и 320i, у Lada одна марка — Vesta. У каждой записи указаны год, цвет, объём двигателя и цена.",
            "На главном экране список выводит характеристики автомобиля. Списки «Производитель» и «Марка» оставляют только нужные записи, кнопка «По цене» меняет порядок сортировки. Кнопка «Добавить» открывает пустую форму. Нажатие на строку открывает ту же форму с данными выбранного автомобиля и сохраняет правку в базу.",
        ],
        "conclusion": "В результате выполнения работы было создано приложение «Авто» со списком автомобилей, добавлением и редактированием записей, фильтрацией по производителю и марке и сортировкой по цене.",
    },
]


def set_run_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, 14, False)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_text(doc, text, *, center=False, bold=False, indent=True, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.first_line_indent = Cm(1.25) if indent and not center else Cm(0)
    run = p.add_run(text)
    set_run_font(run, 14, bold)
    return p


def hide_borders(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        element.set(qn("w:sz"), "0")
        element.set(qn("w:space"), "0")
        borders.append(element)
    tc_pr.append(borders)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(2)
    pf.space_after = Pt(10)
    pf.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, 12, False)
    run.italic = True
    return p


def add_picture(doc, info, lab_no, index):
    path = Path(info["file"])
    if not path.exists():
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    set_run_font(run, 14, False)
    width = Cm(15.2) if info.get("wide") else Cm(6.4)
    run.add_picture(str(path), width=width)
    add_caption(doc, f"Рис. {lab_no}.{index} — {info['caption']}")
    return True


def add_phone_row(doc, pair, lab_no, start_index):
    table = doc.add_table(rows=2, cols=len(pair))
    table.autofit = True
    for col, info in enumerate(pair):
        cell = table.cell(0, col)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].paragraph_format.space_before = Pt(6)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        run = cell.paragraphs[0].add_run()
        set_run_font(run, 12, False)
        run.add_picture(str(info["file"]), width=Cm(6.3))
        cap = table.cell(1, col).paragraphs[0]
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(2)
        cap.paragraph_format.space_after = Pt(8)
        caption = cap.add_run(f"Рис. {lab_no}.{start_index + col} — {info['caption']}")
        set_run_font(caption, 12, False)
        caption.italic = True
        hide_borders(table.cell(0, col))
        hide_borders(table.cell(1, col))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_shots(doc, shots, lab_no):
    index = 1
    cursor = 0
    ready = [info for info in shots if Path(info["file"]).exists()]
    while cursor < len(ready):
        info = ready[cursor]
        if info.get("wide"):
            if add_picture(doc, info, lab_no, index):
                index += 1
            cursor += 1
            continue
        nxt = cursor + 1
        if nxt < len(ready) and not ready[nxt].get("wide"):
            add_phone_row(doc, [info, ready[nxt]], lab_no, index)
            index += 2
            cursor += 2
        else:
            if add_picture(doc, info, lab_no, index):
                index += 1
            cursor += 1


def load_shots(n):
    folder = SHOTS / f"pr{n}"
    meta = folder / "captions.json"
    if not meta.exists():
        return []
    data = json.loads(meta.read_text(encoding="utf-8"))
    return data


SOURCES = [
    "Наугольнов О.А., Елкин Н.С., Январев С.Г., Логанчук М.Л. Разработка мобильных приложений: учебно-методическое пособие для практических работ. Новочеркасск: ЮРГПУ (НПИ), 2024.",
    "Документация Android Developers. URL: https://developer.android.com/studio (дата обращения: 27.09.2026).",
    "Metanit. Руководство по Android. URL: https://metanit.com/java/android/ (дата обращения: 27.09.2026).",
]

SPLIT_NAMES = {
    0: "Отчет_00_Вводное_занятие.docx",
    1: "Отчет_ПР01_Интерфейс_Android_Studio.docx",
    2: "Отчет_ПР02_Поля_ввода.docx",
    3: "Отчет_ПР03_Обработчики_событий.docx",
    4: "Отчет_ПР04_Строковые_и_цветовые_ресурсы.docx",
    5: "Отчет_ПР05_Переключение_между_страницами.docx",
    6: "Отчет_ПР06_Передача_данных.docx",
    7: "Отчет_ПР07_Возврат_результата.docx",
    8: "Отчет_ПР08_Изображения.docx",
    9: "Отчет_ПР09_Аудио.docx",
    10: "Отчет_ПР10_Видео.docx",
    11: "Отчет_ПР11_Покадровая_анимация.docx",
    12: "Отчет_ПР12_Анимация_свойств.docx",
    13: "Отчет_ПР13_Геометрические_фигуры.docx",
    14: "Отчет_ПР14_Авто.docx",
}


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return doc


def add_title_page(doc, lines):
    for line in (
        "Министерство науки и высшего образования Российской Федерации",
        "Южно-Российский государственный политехнический университет (НПИ)",
        "имени М.И. Платова",
    ):
        add_text(doc, line, center=True, indent=False, space_after=0)
    add_text(doc, "ОТЧЁТ", center=True, bold=True, indent=False, space_after=2, space_before=36)
    for line in lines:
        add_text(doc, line, center=True, indent=False, space_after=0)
    add_text(doc, "Выполнил: Яценко Е. А.", center=True, indent=False, space_after=0, space_before=28)
    add_text(doc, "Группа: ________________", center=True, indent=False, space_after=0)
    add_text(doc, "Новочеркасск", center=True, indent=False, space_after=0, space_before=28)
    add_text(doc, "2026", center=True, indent=False)
    doc.add_page_break()


def write_lab(doc, lab):
    prefix = lab["n"] if lab["n"] else "В"
    add_text(doc, lab["title"], bold=True, indent=False, space_after=8, space_before=6)
    add_text(doc, "Цель работы: " + lab["goal"])
    add_text(doc, "Ход работы", bold=True, indent=False, space_before=8)
    for paragraph in lab["body"]:
        add_text(doc, paragraph)
    add_shots(doc, load_shots(lab["n"]), prefix)
    add_text(doc, "Вывод. " + lab["conclusion"], space_before=6)


def add_sources(doc):
    add_text(doc, "Список использованных источников", bold=True, indent=False, space_before=12)
    for i, src in enumerate(SOURCES, start=1):
        add_text(doc, f"{i}. {src}")


def build():
    doc = new_document()
    add_title_page(doc, [
        "по практическим работам",
        "по дисциплине «Разработка мобильных приложений»",
    ])
    add_text(doc, "Цель комплекса работ: изучить технологию создания приложений в среде разработки Android Studio.", indent=True)
    add_text(doc, "Проекты выполнены на языке Java и названы YacenkoEA_PR с номером работы. На экране приложения заголовок имеет вид «Яценко PR» и номер работы. В исходном коде блоки сопровождены пояснениями. Ниже по каждой работе приведены цель, ход выполнения, экранные копии Android Studio и эмулятора и вывод.")
    for lab in LABS:
        doc.add_page_break()
        write_lab(doc, lab)
    doc.add_page_break()
    add_sources(doc)
    doc.save(OUT)
    print(OUT)

    split_dir = ROOT / "Отчеты"
    split_dir.mkdir(exist_ok=True)
    for lab in LABS:
        single = new_document()
        if lab["n"]:
            subtitle = f"по практической работе № {lab['n']}"
        else:
            subtitle = "по вводному занятию"
        add_title_page(single, [
            subtitle,
            "по дисциплине «Разработка мобильных приложений»",
        ])
        write_lab(single, lab)
        add_sources(single)
        path = split_dir / SPLIT_NAMES[lab["n"]]
        single.save(path)
        print(path)


if __name__ == "__main__":
    build()
